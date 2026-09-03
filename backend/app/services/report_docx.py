"""报告（风险评估/应急资源调查）DOCX 公文版式生成器。

复用 docx_template 的预案样式引擎，产出与应急预案一致的版式：
封面（企业名+报告名+日期）→ 页眉页码 → 正文大标题 → 黑体一级标题分页 →
仿宋 16pt 正文 + 规范表格。不含预案“批准页/发布页”。
"""

import io
import logging
import re
from datetime import datetime

import markdown
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from app.services.docx_template import (
    FONT_HEITI,
    MARGIN_COVER_BOTTOM,
    MARGIN_COVER_LEFT,
    MARGIN_COVER_RIGHT,
    MARGIN_COVER_TOP,
    STYLE_COVER_SIGN,
    STYLE_COVER_TITLE,
    add_body_title,
    add_heading,
    add_normal_paragraph,
    add_section,
    html_to_docx_content,
    register_all_styles,
    set_page_margins,
    _set_east_asian_font_in_run,
    _setup_header_footer,
)

logger = logging.getLogger(__name__)

REPORT_KIND_TITLES = {
    "risk": "生产安全事故风险评估报告",
    "resource": "应急资源调查报告",
}


def split_report_content_chapters(content: str) -> list[dict]:
    """按 '## 标题' 切分合并正文，作为 summary.chapters 缺失时的兜底。"""
    chapters: list[dict] = []
    cur_title: str | None = None
    cur_body: list[str] = []

    def flush():
        if cur_title is not None:
            chapters.append({"title": cur_title, "content": "\n".join(cur_body).strip()})

    for line in (content or "").splitlines():
        # 兜底只按二级标题“## 章节”切分；h1 是报告主标题，直接忽略
        m = re.match(r"^\s*##\s+(.+?)\s*$", line)
        if m:
            flush()
            cur_title = m.group(1).strip()
            cur_body = []
            continue
        if cur_title is not None:
            cur_body.append(line)
    flush()
    return chapters


def _build_report_cover(doc: Document, company_name: str, report_label: str):
    """封面：企业名 + 报告名 + 落款日期（不含预案批准页）。"""
    for _ in range(3):
        doc.add_paragraph("")
    doc.add_paragraph(company_name, style=STYLE_COVER_TITLE)
    doc.add_paragraph(report_label, style=STYLE_COVER_TITLE)
    for _ in range(2):
        doc.add_paragraph("")
    sig = doc.add_paragraph(company_name, style=STYLE_COVER_SIGN)
    for run in sig.runs:
        run.font.name = FONT_HEITI
        _set_east_asian_font_in_run(run, FONT_HEITI)
    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = date_p.add_run(datetime.now().strftime("%Y年%m月"))
    r.font.name = FONT_HEITI
    r.font.size = Pt(18)
    _set_east_asian_font_in_run(r, FONT_HEITI)
    doc.add_page_break()


def _clean_chapter(content: str) -> str:
    """复用风险路由清洗：Markdown 管道表→HTML、去 mermaid/尾随 JSON 等残留。"""
    from app.routers.risk_assessment import _clean_for_docx  # 延迟 import 防环依赖
    return _clean_for_docx(content or "")


def _norm_title_text(text: str) -> str:
    """去掉标题行常见的 markdown 修饰（#/**）后归一化。"""
    t = text.strip()
    t = re.sub(r"^#{1,6}\s*", "", t)
    t = t.strip("*").strip()
    return t.strip()


def _strip_leading_title_duplicate(content: str, title: str) -> str:
    """删除章节正文开头与章节标题重复的首行（模型输出常自带标题行）。"""
    target = _norm_title_text(title)
    if not target:
        return content
    lines = (content or "").splitlines()
    for i, line in enumerate(lines):
        stripped = line.strip()
        if not stripped:
            continue
        if _norm_title_text(stripped) == target:
            # 去掉标题行及紧随其后的空行
            rest = lines[i + 1:]
            while rest and not rest[0].strip():
                rest.pop(0)
            return "\n".join(rest)
        break
    return content


def _embed_local_image(doc: Document, src: str) -> bool:
    """把 /uploads/... 本地图嵌入 docx；失败返回 False。"""
    from app.main import UPLOAD_DIR

    path = src
    if src.startswith("/uploads/"):
        path = str(UPLOAD_DIR) + src[len("/uploads"):]
    try:
        with open(path, "rb") as f:
            doc.add_picture(io.BytesIO(f.read()), width=Cm(14.6))
        if doc.paragraphs:
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        return True
    except Exception:
        return False


def generate_report_docx(
    *,
    company_name: str,
    report_kind: str,
    chapters: list[dict],
    report_title: str = "",
) -> Document:
    """生成公文版式报告 DOCX。

    chapters: [{"key", "title", "content"}]，content 为 markdown/HTML 混合。
    """
    report_label = REPORT_KIND_TITLES.get(report_kind, report_title or "调查报告")
    doc = Document()
    register_all_styles(doc)

    first = doc.sections[0]
    set_page_margins(
        first, MARGIN_COVER_LEFT, MARGIN_COVER_RIGHT,
        MARGIN_COVER_TOP, MARGIN_COVER_BOTTOM,
    )
    first.page_width = Cm(21)
    first.page_height = Cm(29.7)

    _build_report_cover(doc, company_name, report_label)
    add_section(
        doc, MARGIN_COVER_LEFT, MARGIN_COVER_RIGHT,
        MARGIN_COVER_TOP, MARGIN_COVER_BOTTOM,
    )
    body_section = doc.sections[-1]
    body_section.header.is_linked_to_previous = True
    body_section.footer.is_linked_to_previous = True
    _setup_header_footer(doc, company_name, report_label)
    add_body_title(doc, report_label)

    for ch in chapters:
        title = (ch.get("title") or "").strip()
        content = (ch.get("content") or "").strip()
        if not title and not content:
            continue
        if title:
            heading = add_heading(doc, title, 1)
            heading.paragraph_format.page_break_before = True
        if not content:
            continue

        cleaned = _strip_leading_title_duplicate(_clean_chapter(content), title)
        html = markdown.markdown(
            cleaned, extensions=["tables", "fenced_code", "md_in_html"],
        )
        # 正文内 <img src="/uploads/..."> 本地嵌入；失败降级为文字说明
        for m in re.finditer(r'<img[^>]*src="(/uploads/[^"]+)"', html):
            if not _embed_local_image(doc, m.group(1)):
                add_normal_paragraph(doc, f"【图片：{m.group(1)} 嵌入失败】")
        html = re.sub(r'<img[^>]*src="/uploads/[^"]+"[^>]*/?>', "", html)
        html_to_docx_content(doc, html, base_level=1)

    return doc
