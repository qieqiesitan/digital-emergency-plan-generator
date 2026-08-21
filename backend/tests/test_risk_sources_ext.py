"""风险源扩展端点回归：模板下载不再 NameError，类别下拉含 27 类。"""
import io

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.database import get_db
from app.dependencies import get_current_user
from app.routers import risk_sources_ext


@pytest.fixture
def client(monkeypatch):
    app = FastAPI()
    app.include_router(risk_sources_ext.router, prefix="/api/v1")

    async def _get_enterprise_data(enterprise_id, user_id, db):
        return {"name": "甲公司", "industry": "工贸", "business_scope": "", "building_overview": "", "employee_count": 10, "address": ""}

    monkeypatch.setattr(risk_sources_ext, "_get_enterprise_data", _get_enterprise_data)

    async def _current_user():
        return type("U", (), {"id": "u1", "email": "admin@test.com"})()

    async def _db():
        yield None

    app.dependency_overrides[get_current_user] = _current_user
    app.dependency_overrides[get_db] = _db
    return TestClient(app)


def test_risk_source_template_uses_2025_categories(client):
    resp = client.get("/api/v1/enterprises/e1/risk-sources/template")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws["A2"].value == "火灾"
    dv = ws.data_validations.dataValidation[0]
    assert "机械致害" in dv.formula1
    assert "锅炉爆炸" not in dv.formula1
