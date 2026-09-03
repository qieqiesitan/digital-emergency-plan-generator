"""报告 docx 公文版式生成器单元测试。"""
import io

from docx import Document

from app.services.report_docx import generate_report_docx


def _doc(kind="risk", chapters=None):
    buf = io.BytesIO()
    doc = generate_report_docx(
        company_name="西安宝岳空间科技有限公司",
        report_kind=kind,
        chapters=chapters
        or [
            {
                "key": "ch1",
                "title": "一、危险有害因素辨识分析",
                "content": (
                    "正文第一段。\n\n**加粗** 内容。\n\n"
                    "| 序号 | 名称 |\n| --- | --- |\n| 1 | 干粉灭火器 |\n"
                    '\n<table border="1"><tr><th>等级</th><th>描述</th></tr>'
                    '<tr><td>低</td><td>可控</td></tr></table>\n'
                    "```mermaid\nflowchart LR\nA-->B\n```\n"
                    '{"overall_assessment": "尾随JSON"}'
                ),
            },
            {"key": "ch2", "title": "二、危险有害因素辨识汇总", "content": "第二章内容。"},
        ],
    )
    doc.save(buf)
    buf.seek(0)
    return Document(buf)


def test_cover_contains_company_and_kind_title_without_approval_page():
    doc = _doc()
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "西安宝岳空间科技有限公司" in joined
    assert "生产安全事故风险评估报告" in joined
    assert "批准页" not in joined


def test_resource_kind_uses_resource_title():
    doc = _doc(kind="resource")
    assert any("应急资源调查报告" in p.text for p in doc.paragraphs)


def test_chapter_headings_present():
    doc = _doc()
    heading_texts = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert "一、危险有害因素辨识分析" in heading_texts
    assert "二、危险有害因素辨识汇总" in heading_texts


def test_clean_removes_mermaid_and_trailing_json():
    doc = _doc()
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "flowchart" not in joined
    assert "尾随JSON" not in joined
    assert "overall_assessment" not in joined


def test_markdown_and_html_tables_become_docx_tables():
    doc = _doc()
    assert len(doc.tables) >= 2
    table_text = "\n".join(
        c.text for t in doc.tables for row in t.rows for c in row.cells
    )
    assert "干粉灭火器" in table_text
    assert "可控" in table_text


def test_first_chapter_heading_has_page_break_before():
    doc = _doc()
    for p in doc.paragraphs:
        if p.style.name == "Heading 1" and "一、" in p.text:
            assert p.paragraph_format.page_break_before is True
            return
    raise AssertionError("未找到分页的一级标题")


def test_builder_uses_chapters_argument():
    doc = _doc(chapters=[{"key": "only", "title": "唯一章节", "content": "只有一章"}])
    assert sum(1 for p in doc.paragraphs if p.style.name == "Heading 1") == 1
