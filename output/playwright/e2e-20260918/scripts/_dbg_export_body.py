"""最小复现：章节正文是否进入预览与 DOCX（调试用，跑完自动清理）。"""
import io
import json
import sys
import urllib.request

from docx import Document

API = "http://localhost:8000/api/v1"
ENT = "10e11995-e682-405a-9035-fbde13cca213"


def call(method, path, token=None, body=None, raw=False):
    data = json.dumps(body).encode() if body is not None else None
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    req = urllib.request.Request(API + path, method=method, data=data, headers=headers)
    with urllib.request.urlopen(req, timeout=120) as resp:
        payload = resp.read()
        return resp.status, (payload if raw else json.loads(payload))


def main():
    _, pr = call("POST", "/auth/login",
                 body={"email": "qa_e2e_test@test.com", "password": "test123456"})
    token = pr["data"]["access_token"]
    _, pl = call("POST", "/plans", token,
                 {"enterprise_id": ENT, "plan_type": "comprehensive", "title": "DBG-导出正文"})
    pid = pl["data"]["id"]
    try:
        _, pr = call("POST", "/auth/login",
                     body={"email": "qa_e2e_test@test.com", "password": "test123456"})
        token = pr["data"]["access_token"]
        _, secs = call("GET", f"/plans/{pid}/sections", token)
        # 复刻完整性探针的写法：25 章全部填满，标题带序号
        filled = 0
        for idx, sec in enumerate(secs["data"], start=1):
            html = (
                f"<h3>{idx} {sec['title']}</h3>"
                f"<p>探针正文：{sec['title']}。本段文字用于验证导出不丢章、不截断。</p>"
            )
            st, _ = call("PUT", f"/plans/{pid}/sections/{sec['section_key']}", token,
                         {"content": html})
            filled += 1 if st == 200 else 0
        print("filled:", filled, "/", len(secs["data"]))
        _, back = call("GET", f"/plans/{pid}/sections/{secs['data'][0]['section_key']}", token)
        print("回读首章:", repr(back["data"]["content"])[:100])

        _, pv = call("GET", f"/plans/{pid}/export/preview", token)
        html = pv["data"]["html"]
        print("preview 含「探针正文」:", "探针正文" in html)
        idx = html.find("总则")
        print("preview 片段:", repr(html[max(0, idx - 200):idx + 300]))

        _, blob = call("POST", f"/plans/{pid}/export/docx", token, raw=True)
        doc = Document(io.BytesIO(blob))
        text = "\n".join(p.text for p in doc.paragraphs)
        print("docx 含「探针正文」:", "探针正文" in text, " paragraphs:", len(doc.paragraphs))
        print("docx 段落 18~48:")
        for i, line in enumerate(text.splitlines()[18:48], start=18):
            print("   ", i, repr(line[:60]))
        print("docx 前 12 段:", [t[:26] for t in text.splitlines()[:12]])
    finally:
        print("delete:", call("DELETE", f"/plans/{pid}", token)[0])
    return 0


if __name__ == "__main__":
    sys.exit(main())
