"""风险源 Excel 导入端到端校验（与资源导入同族的另一条链路）。

同族缺陷刚在资源导入上线过：`Workbook(BytesIO)` 误用导致上传必 500。这条链路同样被修，
这里用真实接口验证「模板 → 预览（含错误行）→ 批量落库 → 台账可查」，并清理。
"""

import io
import json
import os
import subprocess
import sys
import urllib.error
import urllib.request
import uuid

from openpyxl import load_workbook

API = "http://localhost:8000/api/v1"
OUT = os.environ.get("E2E_OUT", "backend/exports/e2e-20260918")
DB = ["docker", "exec", "emergency-plan-db", "psql", "-U", "postgres", "-d", "emergency_plan",
      "-t", "-A", "-q", "-c"]
U, P = "qa_e2e_test@test.com", "test123456"
ENT = "10e11995-e682-405a-9035-fbde13cca213"
NAME_OK, NAME_BAD = "导入探针风险源-甲", "导入探针风险源-坏行"

results = []


def sql(statement: str) -> str:
    # 必须显式 utf-8：Windows 默认 gbk 解码会在中文返回值上炸（本轮踩过）
    out = subprocess.run(DB + [statement], capture_output=True, text=True,
                         encoding="utf-8", errors="replace", check=True)
    for line in out.stdout.splitlines():
        line = line.strip()
        if line and not line.startswith(("INSERT", "DELETE", "UPDATE")):
            return line
    return ""


def check(name, ok, detail=""):
    results.append({"name": name, "ok": bool(ok), "detail": str(detail)[:300]})
    print(("PASS " if ok else "FAIL ") + name + (f" | {detail}" if detail else ""), flush=True)


def call(method, path, token=None, body=None, raw=False):
    data = json.dumps(body).encode() if body is not None else None
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    req = urllib.request.Request(API + path, method=method, data=data, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            payload = resp.read()
            return resp.status, (payload if raw else json.loads(payload))
    except urllib.error.HTTPError as exc:
        return exc.code, exc.read()


def upload(path, token, filename, content: bytes):
    boundary = "----probe" + uuid.uuid4().hex
    body = (
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'
        "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n"
    ).encode() + content + f"\r\n--{boundary}--\r\n".encode()
    req = urllib.request.Request(
        API + path, method="POST", data=body,
        headers={"Authorization": f"Bearer {token}",
                 "Content-Type": f"multipart/form-data; boundary={boundary}"},
    )
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            return resp.status, json.loads(resp.read())
    except urllib.error.HTTPError as exc:
        try:
            return exc.code, json.loads(exc.read())
        except Exception:
            return exc.code, {}


def main():
    try:
        _, pr = call("POST", "/auth/login", body={"email": U, "password": P})
        token = pr["data"]["access_token"]
        status, tpl = call("GET", f"/enterprises/{ENT}/risk-sources/template", token, raw=True)
        check("下载风险源模板", status == 200 and isinstance(tpl, bytes) and len(tpl) > 2000,
              f"status={status} bytes={len(tpl) if isinstance(tpl, bytes) else '-'}")
        wb = load_workbook(io.BytesIO(tpl))
        ws = wb.active
        headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
        check("模板表头符合契约",
              headers[:4] == ["风险类别", "风险名称", "位置", "风险描述"], headers)
        ws.append(["火灾", NAME_OK, "一号库房", "电气线路老化", "中", "高", "定期检测"])
        ws.append(["火灾", NAME_BAD, "", "", "极高", "极高", ""])   # 可能性/严重性非法值
        buf = io.BytesIO()
        wb.save(buf)

        status, res = upload(f"/enterprises/{ENT}/risk-sources/import", token,
                             "risk_sources.xlsx", buf.getvalue())
        data = res.get("data", {}) if isinstance(res, dict) else {}
        check("导入预览返回 200（不再 500）", status == 200, f"status={status}")
        check("预览统计正确（1 有效 / 1 错误 / 模板示例行跳过）",
              data.get("valid_count") == 1 and data.get("error_count") == 1
              and data.get("skipped_examples") == 1,
              f"valid={data.get('valid_count')} error={data.get('error_count')} "
              f"skipped={data.get('skipped_examples')}")
        check("模板示例行未进入预览", all(
            "【示例】" not in (i["data"].get("name") or "") for i in (data.get("items") or [])))
        ok_items = [i for i in (data.get("items") or [])
                    if not i.get("errors") and i["data"]["name"] == NAME_OK]
        check("有效行内容正确", bool(ok_items), ok_items)
        check("预览不改库",
              sql(f"select count(*) from risk_sources where enterprise_id='{ENT}' "
                  f"and name='{NAME_OK}';") == "0")

        status, created = call("POST", f"/enterprises/{ENT}/risk-sources/batch", token,
                               {"items": [ok_items[0]["data"]]})
        check("批量落库返回 201", status == 201, f"status={status}")
        check("台账可查到风险源", sql(
            f"select count(*) from risk_sources where enterprise_id='{ENT}' "
            f"and name='{NAME_OK}';") == "1")
        level = sql(f"select coalesce(risk_level,'') from risk_sources where enterprise_id='{ENT}' "
                    f"and name='{NAME_OK}';")
        check("落库时按 可能性×严重性 算了风险等级", bool(level), f"risk_level={level}")
    finally:
        sql(f"delete from risk_sources where enterprise_id='{ENT}' "
            f"and name in ('{NAME_OK}','{NAME_BAD}');")
        print("cleanup done: 探针风险源已删除")

    passed = sum(1 for r in results if r["ok"])
    print(f"合计 {passed}/{len(results)} PASS")
    with open(os.path.join(OUT, "summary-risk-source-import.json"), "w", encoding="utf-8") as fh:
        json.dump({"passed": passed, "total": len(results), "results": results}, fh,
                  ensure_ascii=False, indent=2)
    return 0 if passed == len(results) else 1


if __name__ == "__main__":
    sys.exit(main())
