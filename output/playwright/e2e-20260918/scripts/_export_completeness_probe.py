"""完整预案导出端到端校验（对应用户核心诉求："要能导出企业的完整预案"）。

流程：新建一张综合预案 → 把模板生成的**全部章节**写满正文（其中一章内嵌 Mermaid 图）
→ 调用导出前校验 → 导出 DOCX → 用 python-docx 逐项核对：
  1. 每个章节标题都出现在文档里（不丢章）；
  2. 正文关键词都在（不截断）；
  3. 含图章节产出内嵌图片（Mermaid 渲染链路通）；
  4. 全文无占位残留（待补充/XXX/{{...}} 等）；
  5. 导出前校验接口对"填满的预案"判定为 valid。
跑完删除探针预案（含章节）。
"""

import io
import json
import os
import re
import sys
import urllib.error
import urllib.request

from docx import Document

API = "http://localhost:8000/api/v1"
OUT = os.environ.get("E2E_OUT", "backend/exports/e2e-20260918")
U, P = "qa_e2e_test@test.com", "test123456"
ENT = "10e11995-e682-405a-9035-fbde13cca213"
PLAN_TITLE = "导出完整性探针（自动清理）"
PLACEHOLDERS = ["待补充", "XXX", "{{", "占位", "TODO", "undefined", "None"]

results = []
token_holder: dict = {}


def check(name, ok, detail=""):
    results.append({"name": name, "ok": bool(ok), "detail": str(detail)[:300]})
    print(("PASS " if ok else "FAIL ") + name + (f" | {detail}" if detail else ""), flush=True)


def call(method, path, token=None, body=None, raw=False):
    data = json.dumps(body).encode() if body is not None else None
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    req = urllib.request.Request(API + path, method=method, data=data, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=180) as resp:
            payload = resp.read()
            return resp.status, (payload if raw else json.loads(payload)), dict(resp.headers)
    except urllib.error.HTTPError as exc:
        payload = exc.read()
        try:
            return exc.code, (payload if raw else json.loads(payload)), dict(exc.headers)
        except Exception:
            return exc.code, payload, dict(exc.headers)


def norm(text: str) -> str:
    return re.sub(r"[\s　]+", "", text)


def main():
    os.makedirs(OUT, exist_ok=True)
    plan_id = None
    try:
        status, payload, _ = call("POST", "/auth/login", body={"email": U, "password": P})
        assert status == 200, payload
        token = payload["data"]["access_token"]
        token_holder["t"] = token

        status, payload, _ = call(
            "POST", "/plans", token=token,
            body={"enterprise_id": ENT, "plan_type": "comprehensive", "title": PLAN_TITLE},
        )
        assert status == 201, (status, payload)
        plan_id = payload["data"]["id"]
        print("plan:", plan_id)

        status, payload, _ = call("GET", f"/plans/{plan_id}/sections", token=token)
        assert status == 200, payload
        sections = payload["data"]
        check("模板生成章节", len(sections) > 0, f"{len(sections)} 章")

        filled = 0
        diagram_key = None
        for idx, sec in enumerate(sections, start=1):
            key, title = sec["section_key"], sec["title"]
            # 组织体系章节内嵌一张 Mermaid 图（只放第一处命中的章节，保持证据确定）
            if diagram_key is None and key in ("sec_3", "sec_3_1"):
                diagram_key = key
            body_html = (
                f"<h3>{idx} {title}</h3>"
                f"<p>探针正文：{title}。本段文字用于验证导出不丢章、不截断。</p>"
            )
            if key == diagram_key:
                body_html += (
                    "<p>组织架构如下：</p>"
                    '<pre><code class="language-mermaid">'
                    "graph TD; A[应急指挥部]-->B[抢险救援组]; A-->C[医疗救护组];"
                    "</code></pre>"
                )
            status, resp, _ = call(
                "PUT", f"/plans/{plan_id}/sections/{key}", token=token, body={"content": body_html}
            )
            if status == 200:
                filled += 1
        check("全部章节写入成功", filled == len(sections), f"{filled}/{len(sections)}")

        status, validate, _ = call("POST", f"/plans/{plan_id}/export/validate", token=token)
        check("导出前校验接口可用", status == 200, f"status={status}")
        vdata = validate.get("data", {}) if isinstance(validate, dict) else {}
        check("填满的预案通过导出前校验", bool(vdata.get("valid")),
              f"issues={vdata.get('issues')} warnings={vdata.get('warnings')}")

        status, blob, headers = call("POST", f"/plans/{plan_id}/export/docx", token=token, raw=True)
        check("DOCX 导出返回 200", status == 200, f"status={status} size={len(blob) if isinstance(blob, bytes) else '-'}")
        if status != 200 or not isinstance(blob, bytes):
            raise SystemExit(1)
        docx_path = os.path.join(OUT, "export-completeness-probe.docx")
        with open(docx_path, "wb") as fh:
            fh.write(blob)

        doc = Document(io.BytesIO(blob))
        text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                text += "\n" + " ".join(c.text for c in row.cells)
        flat = norm(text)

        missing = [s["title"] for s in sections if norm(s["title"]).replace("　", "") not in flat]
        check("每个章节标题都出现在导出文档里", not missing,
              f"缺失 {len(missing)}：{missing[:5]}")

        missing_body = [
            s["title"] for s in sections if f"探针正文：{s['title']}" not in text
        ]
        check("每章正文都在（无截断/丢内容）", not missing_body,
              f"缺正文 {len(missing_body)}：{missing_body[:5]}")

        images = len(doc.inline_shapes)
        check("含图章节产出内嵌图片", images >= 1, f"inline_images={images}")

        hits = sorted({ph for ph in PLACEHOLDERS if ph in text})
        check("全文无占位符残留", not hits, f"命中={hits}")

        check("文档体量合理（非空壳）", len(blob) > 20_000, f"{len(blob)} bytes")
        print("docx:", docx_path)
    finally:
        if plan_id:
            try:
                status, _, _ = call("DELETE", f"/plans/{plan_id}", token=token_holder.get("t"))
                print(f"cleanup: delete plan -> {status}")
            except Exception as exc:  # noqa: BLE001
                print("CLEANUP FAILED:", exc)

    passed = sum(1 for r in results if r["ok"])
    print(f"合计 {passed}/{len(results)} PASS")
    with open(os.path.join(OUT, "summary-export-completeness.json"), "w", encoding="utf-8") as fh:
        json.dump({"passed": passed, "total": len(results), "results": results},
                  fh, ensure_ascii=False, indent=2)
    return 0 if passed == len(results) else 1


if __name__ == "__main__":
    sys.exit(main())
