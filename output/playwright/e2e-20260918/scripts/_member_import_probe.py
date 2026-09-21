"""成员 Excel 批量导入端到端校验：下载模板 → 填行 → 上传 → 核对入库/建节点/错误行。

这张表是给企业 HR 用的（一次导几百人），最容易在重构后静默坏掉；
探针自带清理：删除导入的成员并把组织树恢复成导入前快照。
"""

import io
import json
import os
import subprocess
import sys
import urllib.error
import urllib.request
import uuid

from openpyxl import load_workbook

API = "http://localhost:8000/api/v1"
OUT = os.environ.get("E2E_OUT", "backend/exports/e2e-20260918")
DB = ["docker", "exec", "emergency-plan-db", "psql", "-U", "postgres", "-d", "emergency_plan",
      "-t", "-A", "-q", "-c"]
U, P = "qa_e2e_test@test.com", "test123456"
ENT = "10e11995-e682-405a-9035-fbde13cca213"
NAME_A, NAME_B, NAME_BAD = "导入探针甲", "导入探针乙", "导入探针丙"

results = []


def sql(statement: str) -> str:
    out = subprocess.run(DB + [statement], capture_output=True, text=True, check=True)
    for line in out.stdout.splitlines():
        line = line.strip()
        if line and not line.startswith(("INSERT", "DELETE", "UPDATE")):
            return line
    return ""


def check(name, ok, detail=""):
    results.append({"name": name, "ok": bool(ok), "detail": str(detail)[:300]})
    print(("PASS " if ok else "FAIL ") + name + (f" | {detail}" if detail else ""), flush=True)


def call(method, path, token=None, body=None, raw=False):
    data = json.dumps(body).encode() if body is not None else None
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    req = urllib.request.Request(API + path, method=method, data=data, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            payload = resp.read()
            return resp.status, (payload if raw else json.loads(payload))
    except urllib.error.HTTPError as exc:
        return exc.code, exc.read()


def upload(path, token, filename, content: bytes):
    boundary = "----probe" + uuid.uuid4().hex
    body = (
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'
        "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n"
    ).encode() + content + f"\r\n--{boundary}--\r\n".encode()
    req = urllib.request.Request(
        API + path, method="POST", data=body,
        headers={"Authorization": f"Bearer {token}",
                 "Content-Type": f"multipart/form-data; boundary={boundary}"},
    )
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            return resp.status, json.loads(resp.read())
    except urllib.error.HTTPError as exc:
        try:
            return exc.code, json.loads(exc.read())
        except Exception:
            return exc.code, {}


def main():
    os.makedirs(OUT, exist_ok=True)
    snapshot = None
    try:
        status, pr = call("POST", "/auth/login", body={"email": U, "password": P})
        token = pr["data"]["access_token"]
        snapshot = sql(
            f"select coalesce(org_structure::text,'[]') from enterprises where id='{ENT}';"
        )

        status, tpl = call("GET", f"/enterprises/{ENT}/org/members/template", token, raw=True)
        check("下载导入模板", status == 200 and isinstance(tpl, bytes) and len(tpl) > 2000,
              f"status={status} bytes={len(tpl) if isinstance(tpl, bytes) else '-'}")
        wb = load_workbook(io.BytesIO(tpl))
        ws = wb.active
        headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
        check("模板表头符合契约", headers[:6] == ["姓名", "邮箱", "部门", "班组", "岗位", "角色"], headers)

        rows = [
            [NAME_A, "", "导入探针生产部", "一班", "操作工", "成员"],
            [NAME_B, "", "导入探针生产部", "二班", "班长", "班组长"],
            [NAME_BAD, "no-such-user@test.invalid", "导入探针生产部", "", "员工", "成员"],
            ["", "", "导入探针生产部", "", "", "成员"],
        ]
        for row in rows:
            ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)

        status, res = upload(f"/enterprises/{ENT}/org/members/import", token,
                             "member_import.xlsx", buf.getvalue())
        data = res.get("data", {}) if isinstance(res, dict) else {}
        check("导入接口返回 200", status == 200, f"status={status}")
        check("成功导入 2 行", data.get("imported") == 2, data)
        errs = data.get("errors") or []
        check("错误行被准确报出（不存在用户 + 姓名为空）",
              len(errs) == 2 and any("不存在" in str(e.get("reason")) for e in errs)
              and any("姓名" in str(e.get("reason")) for e in errs), errs)

        status, members = call("GET", f"/enterprises/{ENT}/org/members", token)
        names = [m.get("name") for m in (members.get("data") or [])]
        check("两位成员已入库", NAME_A in names and NAME_B in names, names)
        status, nodes = call("GET", f"/enterprises/{ENT}/org/nodes", token)
        node_names = [n.get("name") for n in (nodes.get("data") or [])]
        check("部门/班组节点按需创建",
              "导入探针生产部" in node_names and "一班" in node_names and "二班" in node_names,
              node_names)
        bad = [m for m in (members.get("data") or []) if m.get("name") == NAME_BAD]
        check("错误行未入库", not bad, bad)
    finally:
        try:
            sql(f"delete from enterprise_members where enterprise_id='{ENT}' "
                f"and name in ('{NAME_A}','{NAME_B}','{NAME_BAD}');")
            if snapshot is not None:
                escaped = snapshot.replace("'", "''")
                sql(f"update enterprises set org_structure='{escaped}'::jsonb where id='{ENT}';")
            print("cleanup done: 成员与组织树已还原")
        except Exception as exc:  # noqa: BLE001
            print("CLEANUP FAILED:", exc)

    passed = sum(1 for r in results if r["ok"])
    print(f"合计 {passed}/{len(results)} PASS")
    with open(os.path.join(OUT, "summary-member-import.json"), "w", encoding="utf-8") as fh:
        json.dump({"passed": passed, "total": len(results), "results": results}, fh,
                  ensure_ascii=False, indent=2)
    return 0 if passed == len(results) else 1


if __name__ == "__main__":
    sys.exit(main())
