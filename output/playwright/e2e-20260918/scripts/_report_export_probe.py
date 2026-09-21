"""报告类导出校验（风险评估 / 资源调查共用 report_docx 渲染器）。

做法：为 QA 企业插入一条合成报告（内容含二级标题章节、以及"正文里提到章节名"的短段落，
用来顺带验证是否存在与 N-13 同类的误删），调用导出接口下载 DOCX 并逐章核对，
最后删除合成报告行。
"""

import io
import json
import os
import re
import subprocess
import sys
import urllib.error
import urllib.request

from docx import Document

API = "http://localhost:8000/api/v1"
OUT = os.environ.get("E2E_OUT", "backend/exports/e2e-20260918")
DB = ["docker", "exec", "emergency-plan-db", "psql", "-U", "postgres", "-d", "emergency_plan",
      "-t", "-A", "-q", "-c"]
U, P = "qa_e2e_test@test.com", "test123456"
ENT = "10e11995-e682-405a-9035-fbde13cca213"
KIND = os.environ.get("E2E_KIND", "risk")   # risk | resource
KINDS = {
    "risk": ("risk_assessment_reports", "/risk-assessment/export", "探针风险评估报告",
             "export-risk-report.docx"),
    "resource": ("resource_investigation_reports", "/resource-investigation/export",
                 "探针资源调查报告", "export-resource-report.docx"),
}

CHAPTERS = [
    ("第一章 评估目的", "本章说明评估目的。为落实安全生产主体责任，特开展本次风险评估。"),
    ("第二章 评估依据", "依据《中华人民共和国安全生产法》及 GB/T 29639-2020 开展评估。"),
    ("第三章 风险辨识与分级", "经辨识，罐区构成较大风险；生产车间构成一般风险。"),
    ("第四章 结论与建议", "综上，本企业风险总体可控，建议持续完善管控措施并定期复评。"),
]
CONTENT = "\n\n".join(f"## {title}\n\n{body}" for title, body in CHAPTERS)


def sql(statement: str) -> str:
    out = subprocess.run(DB + [statement], capture_output=True, text=True, check=True)
    for line in out.stdout.splitlines():
        line = line.strip()
        if line and not line.startswith(("INSERT", "DELETE", "UPDATE")):
            return line
    return ""


def call(method, path, token=None, body=None, raw=False):
    data = json.dumps(body).encode() if body is not None else None
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    req = urllib.request.Request(API + path, method=method, data=data, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=180) as resp:
            payload = resp.read()
            return resp.status, (payload if raw else json.loads(payload))
    except urllib.error.HTTPError as exc:
        return exc.code, exc.read()


def main():
    os.makedirs(OUT, exist_ok=True)
    report_id = None
    try:
        _, pr = call("POST", "/auth/login", body={"email": U, "password": P})
        token = pr["data"]["access_token"]
        table, endpoint, title_text, out_name = KINDS[KIND]
        report_id = sql(
            f"insert into {table} "
            "(id, enterprise_id, title, content, summary, status, generated_by) values "
            f"(gen_random_uuid(),'{ENT}','{title_text}', $doc${CONTENT}$doc$, '{{}}'::jsonb,"
            " 'completed','ai') returning id;"
        )
        print(f"[{KIND}] report:", report_id)
        status, blob = call("GET", f"/enterprises/{ENT}{endpoint}", token, raw=True)
        print("export status:", status, "size:", len(blob) if isinstance(blob, bytes) else "-")
        if status != 200 or not isinstance(blob, bytes):
            print("body:", blob[:200])
            return 1
        path = os.path.join(OUT, out_name)
        with open(path, "wb") as fh:
            fh.write(blob)
        doc = Document(io.BytesIO(blob))
        text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                text += "\n" + " ".join(c.text for c in row.cells)
        flat = re.sub(r"\s+", "", text)

        fails = []
        for title, body in CHAPTERS:
            hit_title = re.sub(r"\s+", "", title) in flat
            hit_body = re.sub(r"\s+", "", body)[:24] in flat
            print(("PASS " if hit_title and hit_body else "FAIL ")
                  + f"{title} | 标题={hit_title} 正文={hit_body}")
            if not (hit_title and hit_body):
                fails.append(title)
        print("段落数:", len(doc.paragraphs), " inline_images:", len(doc.inline_shapes))
        with open(os.path.join(OUT, f"summary-report-export-{KIND}.json"), "w", encoding="utf-8") as fh:
            json.dump({"kind": KIND, "status": status, "size": len(blob), "fails": fails,
                       "paragraphs": len(doc.paragraphs)}, fh, ensure_ascii=False, indent=2)
        return 0 if not fails else 1
    finally:
        if report_id:
            sql(f"delete from {KINDS[KIND][0]} where id='{report_id}';")
            print("cleanup done")


if __name__ == "__main__":
    sys.exit(main())
