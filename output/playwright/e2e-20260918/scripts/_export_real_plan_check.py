"""真实内容预案导出校验（只读）：用 QA 名下已有的、含 AI 生成正文的预案导出 DOCX，
逐章核对"正文片段是否出现在文档里"，作为合成探针之外的现实样本验证。

不写任何数据：只 GET 章节 + POST 导出。
"""

import io
import json
import os
import re
import sys
import urllib.request

from docx import Document

API = "http://localhost:8000/api/v1"
OUT = os.environ.get("E2E_OUT", "backend/exports/e2e-20260918")
PLAN = os.environ.get("E2E_PLAN", "aa8244ab-ff8f-45a4-8814-15629041095f")
U, P = "qa_e2e_test@test.com", "test123456"


def call(method, path, token=None, body=None, raw=False):
    data = json.dumps(body).encode() if body is not None else None
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    req = urllib.request.Request(API + path, method=method, data=data, headers=headers)
    with urllib.request.urlopen(req, timeout=180) as resp:
        payload = resp.read()
        return resp.status, (payload if raw else json.loads(payload))


def plain(html_text: str) -> str:
    text = re.sub(r"<[^>]+>", " ", html_text or "")
    return re.sub(r"\s+", "", text)


def main():
    os.makedirs(OUT, exist_ok=True)
    _, pr = call("POST", "/auth/login", body={"email": U, "password": P})
    token = pr["data"]["access_token"]
    _, secs = call("GET", f"/plans/{PLAN}/sections", token)
    filled = [s for s in secs["data"] if (s.get("content") or "").strip()]
    print(f"预案 {PLAN}: 非空章节 {len(filled)}/{len(secs['data'])}")

    _, blob = call("POST", f"/plans/{PLAN}/export/docx", token, raw=True)
    path = os.path.join(OUT, "export-real-plan-check.docx")
    with open(path, "wb") as fh:
        fh.write(blob)
    doc = Document(io.BytesIO(blob))
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            text += "\n" + " ".join(c.text for c in row.cells)
    flat = plain(text)

    missing = []
    for sec in filled:
        # 指纹取自「最长的那个段落」：章节首段常是重复标题（会被导出正确剥离），
        # 从整段内容里取指纹会跨过标题边界导致误判（本探针第一版就踩了）。
        paragraphs = [plain(p) for p in re.split(r"</(?:p|div|li|h[1-6])>", sec["content"] or "")]
        paragraphs = [p for p in paragraphs if len(p) >= 20]
        body = max(paragraphs, key=len) if paragraphs else plain(sec["content"])
        probe = body[:40]
        hit = probe in flat
        print(("PASS " if hit else "FAIL ") + f"{sec['section_key']} {sec['title'][:12]} 指纹命中={hit}")
        if not hit:
            missing.append((sec["section_key"], probe[:30]))

    print("inline_images =", len(doc.inline_shapes), " size =", len(blob))
    print("未命中的章节:", missing)
    ok = not missing
    with open(os.path.join(OUT, "summary-export-real-plan.json"), "w", encoding="utf-8") as fh:
        json.dump({"plan": PLAN, "filled_sections": len(filled), "missing": missing,
                   "inline_images": len(doc.inline_shapes), "size": len(blob), "ok": ok},
                  fh, ensure_ascii=False, indent=2)
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
