"""四色图工作台端到端校验：楼层→分区→风险对象→单元→事件→措施，
再验工作台读取、管控清单（含 Excel 导出）、风险公示 token、转换参考，最后验删除级联。

全部零 AI：risk-management 的 CRUD/计算/导出都不依赖模型（只有 /ai/* 才需要）。
探针自带清理：按创建顺序逆序删除，且校验删除楼层后子数据确实级联清掉。
"""

import io
import json
import os
import subprocess
import sys
import urllib.error
import urllib.request

from openpyxl import load_workbook

API = "http://localhost:8000/api/v1"
OUT = os.environ.get("E2E_OUT", "backend/exports/e2e-20260918")
DB = ["docker", "exec", "emergency-plan-db", "psql", "-U", "postgres", "-d", "emergency_plan",
      "-t", "-A", "-q", "-c"]
U, P = "qa_e2e_test@test.com", "test123456"
ENT = "10e11995-e682-405a-9035-fbde13cca213"
BASE = f"/enterprises/{ENT}/risk-management"
PFX = "探针-四色图"

results = []
created = {}


def sql(statement: str) -> str:
    out = subprocess.run(DB + [statement], capture_output=True, text=True,
                         encoding="utf-8", errors="replace", check=True)
    for line in out.stdout.splitlines():
        line = line.strip()
        if line and not line.startswith(("INSERT", "DELETE", "UPDATE")):
            return line
    return ""


def check(name, ok, detail=""):
    results.append({"name": name, "ok": bool(ok), "pass": bool(ok), "detail": str(detail)[:300]})
    print(("PASS " if ok else "FAIL ") + name + (f" | {detail}" if detail else ""), flush=True)


def call(method, path, token=None, body=None, raw=False, expect=None, label=""):
    data = json.dumps(body).encode() if body is not None else None
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    req = urllib.request.Request(API + path, method=method, data=data, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=180) as resp:
            status, raw_bytes = resp.status, resp.read()
    except urllib.error.HTTPError as exc:
        status, raw_bytes = exc.code, exc.read()
    if raw:
        payload = raw_bytes
    else:
        try:
            payload = json.loads(raw_bytes)
        except Exception:
            payload = raw_bytes[:200].decode("utf-8", "replace")
    if expect is not None:
        ok = status == expect
        results.append({"label": label, "expect": expect, "got": status, "pass": ok})
        print(("PASS " if ok else "FAIL ") + f"{label}（期望 {expect}，实际 {status}）")
    return status, payload


def main():
    os.makedirs(OUT, exist_ok=True)
    try:
        _, pr = call("POST", "/auth/login", body={"email": U, "password": P})
        token = pr["data"]["access_token"]

        status, floor = call("POST", f"{BASE}/floors?enterprise_id={ENT}", token,
                             {"name": f"{PFX}-楼层", "sort_order": 98}, expect=201,
                             label="建楼层")
        created["floor"] = (floor.get("data") or {}).get("id")
        status, zone = call("POST", f"{BASE}/zones?enterprise_id={ENT}", token,
                            {"name": f"{PFX}-分区", "floor_id": created["floor"]}, expect=201,
                            label="建分区")
        created["zone"] = (zone.get("data") or {}).get("id")
        status, obj = call("POST", f"{BASE}/objects?enterprise_id={ENT}", token,
                           {"name": f"{PFX}-风险点", "zone_id": created["zone"],
                            "is_risk_point": True, "location_x": 12.5, "location_y": 33.0,
                            "responsible_person": "探针责任人", "contact_phone": "13800000000"},
                           expect=201, label="建风险对象（风险点，带坐标）")
        created["object"] = (obj.get("data") or {}).get("id")
        status, unit = call("POST", f"{BASE}/objects/{created['object']}/units?enterprise_id={ENT}",
                            token, {"name": f"{PFX}-单元", "unit_type": "设备"},
                            expect=201, label="建单元")
        created["unit"] = (unit.get("data") or {}).get("id")
        status, event = call("POST", f"{BASE}/units/{created['unit']}/events?enterprise_id={ENT}",
                             token, {"accident_type": "火灾", "description": "探针事件",
                                     "method_type": "LS", "method_params": {"l": 3, "s": 4},
                                     "inherent_risk_level": "重大"}, expect=201, label="建事件（LS 法）")
        created["event"] = (event.get("data") or {}).get("id")

        status, recalc = call("POST", f"{BASE}/events/{created['event']}/recalc?enterprise_id={ENT}",
                              token, expect=200, label="事件重算")
        data = recalc.get("data") or {}
        check("重算分值与等级正确（L=3,S=4 → R=12）",
              data.get("risk_score") == "R=12" and bool(data.get("risk_level")),
              f"score={data.get('risk_score')} level={data.get('risk_level')}")

        status, _ = call("POST", f"{BASE}/events/{created['event']}/measures?enterprise_id={ENT}",
                         token, {"measure_category": "工程技术", "description": "探针措施：加装报警",
                                 "responsible_person": "探针责任人"}, expect=201, label="建措施")
        status, measures = call("GET",
                                f"{BASE}/events/{created['event']}/measures?enterprise_id={ENT}",
                                token)
        check("措施列表可读且含 1 条",
              status == 200 and len(measures.get("data") or []) == 1, f"status={status}")

        status, wb = call("GET", f"{BASE}/workbench?enterprise_id={ENT}&floor_id={created['floor']}",
                          token)
        zones = (wb.get("data") or {}).get("zones") if isinstance(wb, dict) else None
        check("工作台返回该楼层分区", status == 200 and isinstance(zones, list) and len(zones) >= 1,
              f"status={status} zones={len(zones) if isinstance(zones, list) else zones}")

        status, hierarchy = call("GET", f"{BASE}/hierarchy?enterprise_id={ENT}", token)
        check("层级树可读", status == 200, f"status={status}")

        # 清单缺省按"企业默认楼层"取值；探针楼不是默认楼，必须显式传 floor_id
        status, clist = call(
            "GET", f"{BASE}/control-list?enterprise_id={ENT}&floor_id={created['floor']}", token)
        text = json.dumps(clist, ensure_ascii=False)
        check("管控清单包含探针风险点", status == 200 and f"{PFX}-风险点" in text,
              f"status={status}")

        status, xlsx = call(
            "GET", f"{BASE}/control-list/export?enterprise_id={ENT}&floor_id={created['floor']}",
            token, raw=True)
        ok = status == 200 and isinstance(xlsx, bytes) and xlsx[:2] == b"PK"
        rows = 0
        if ok:
            sheet = load_workbook(io.BytesIO(xlsx)).active
            rows = sheet.max_row
        check("管控清单 Excel 导出可打开且有数据行", ok and rows >= 2,
              f"status={status} bytes={len(xlsx) if isinstance(xlsx, bytes) else '-'} rows={rows}")

        status, tok = call("POST", f"{BASE}/risk-publicity/token?enterprise_id={ENT}", token)
        check("风险公示 token 可生成", status == 200, f"status={status}")

        status, conv = call("GET",
                            f"{BASE}/events/{created['event']}/conversion-reference?enterprise_id={ENT}",
                            token)
        check("事件转换参考可读", status == 200, f"status={status}")

        # ---- 删除级联：删楼层后子数据应一并消失 ----
        status, deleted = call("DELETE", f"{BASE}/floors/{created['floor']}?enterprise_id={ENT}",
                               token, expect=200, label="删除楼层")
        left = {
            "zone": sql(f"select count(*) from risk_zones where id='{created['zone']}';"),
            "object": sql(f"select count(*) from risk_objects where id='{created['object']}';"),
            "unit": sql(f"select count(*) from risk_units where id='{created['unit']}';"),
            "event": sql(f"select count(*) from risk_events where id='{created['event']}';"),
        }
        check("删除楼层级联清掉分区/对象/单元/事件",
              all(v == "0" for v in left.values()), left)
    finally:
        # 兜底清理（级联失效时也能收拾干净）
        if created.get("event"):
            sql(f"delete from risk_measures where event_id='{created['event']}';")
            sql(f"delete from risk_events where id='{created['event']}';")
        if created.get("unit"):
            sql(f"delete from risk_units where id='{created['unit']}';")
        if created.get("object"):
            sql(f"delete from risk_objects where id='{created['object']}';")
        if created.get("zone"):
            sql(f"delete from risk_zones where id='{created['zone']}';")
        if created.get("floor"):
            sql(f"delete from enterprise_floors where id='{created['floor']}';")
        print("cleanup done: 探针四色图数据已清理")

    passed = sum(1 for r in results if r.get("ok") or r.get("pass"))
    print(f"合计 {passed}/{len(results)} PASS")
    with open(os.path.join(OUT, "summary-risk-workbench.json"), "w", encoding="utf-8") as fh:
        json.dump({"passed": passed, "total": len(results), "results": results}, fh,
                  ensure_ascii=False, indent=2)
    return 0 if passed == len(results) else 1


if __name__ == "__main__":
    sys.exit(main())
