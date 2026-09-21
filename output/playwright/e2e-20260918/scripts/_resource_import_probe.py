"""应急资源 Excel 导入端到端校验：模板 → 预览（含错误行）→ 批量落库 → 台账可见 → 清理。

资源台账是"完整预案"里附件与资源配置章节的数据来源，导入坏了会一路影响导出内容。
"""

import io
import json
import os
import subprocess
import sys
import urllib.error
import urllib.request
import uuid

from openpyxl import load_workbook

API = "http://localhost:8000/api/v1"
OUT = os.environ.get("E2E_OUT", "backend/exports/e2e-20260918")
DB = ["docker", "exec", "emergency-plan-db", "psql", "-U", "postgres", "-d", "emergency_plan",
      "-t", "-A", "-q", "-c"]
U, P = "qa_e2e_test@test.com", "test123456"
ENT = "10e11995-e682-405a-9035-fbde13cca213"
NAME_OK, NAME_BAD = "导入探针-干粉灭火器", "导入探针-未知类别"

results = []


def sql(statement: str) -> str:
    out = subprocess.run(DB + [statement], capture_output=True, text=True, check=True)
    for line in out.stdout.splitlines():
        line = line.strip()
        if line and not line.startswith(("INSERT", "DELETE", "UPDATE")):
            return line
    return ""


def check(name, ok, detail=""):
    results.append({"name": name, "ok": bool(ok), "detail": str(detail)[:300]})
    print(("PASS " if ok else "FAIL ") + name + (f" | {detail}" if detail else ""), flush=True)


def call(method, path, token=None, body=None, raw=False):
    data = json.dumps(body).encode() if body is not None else None
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    req = urllib.request.Request(API + path, method=method, data=data, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            payload = resp.read()
            return resp.status, (payload if raw else json.loads(payload))
    except urllib.error.HTTPError as exc:
        return exc.code, exc.read()


def upload(path, token, filename, content: bytes):
    boundary = "----probe" + uuid.uuid4().hex
    body = (
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'
        "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n"
    ).encode() + content + f"\r\n--{boundary}--\r\n".encode()
    req = urllib.request.Request(
        API + path, method="POST", data=body,
        headers={"Authorization": f"Bearer {token}",
                 "Content-Type": f"multipart/form-data; boundary={boundary}"},
    )
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            return resp.status, json.loads(resp.read())
    except urllib.error.HTTPError as exc:
        try:
            return exc.code, json.loads(exc.read())
        except Exception:
            return exc.code, {}


def main():
    os.makedirs(OUT, exist_ok=True)
    try:
        _, pr = call("POST", "/auth/login", body={"email": U, "password": P})
        token = pr["data"]["access_token"]
        status, tpl = call("GET", f"/enterprises/{ENT}/resources/template", token, raw=True)
        check("下载资源导入模板", status == 200 and isinstance(tpl, bytes) and len(tpl) > 2000,
              f"status={status} bytes={len(tpl) if isinstance(tpl, bytes) else '-'}")
        wb = load_workbook(io.BytesIO(tpl))
        ws = wb.active
        headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
        check("模板表头符合契约",
              headers[:5] == ["类别", "名称", "规格型号", "数量", "单位"], headers)
        ws.append(["消防设施", NAME_OK, "MFZ/ABC4", 12, "具", "一号库房", "王主管", "13800000000", "否", "", ""])
        ws.append([NAME_BAD, "坏行资源", "", 1, "件", "", "", "", "否", "", ""])
        buf = io.BytesIO()
        wb.save(buf)

        status, res = upload(f"/enterprises/{ENT}/resources/import", token,
                             "resource_import.xlsx", buf.getvalue())
        data = res.get("data", {}) if isinstance(res, dict) else {}
        check("导入预览返回 200", status == 200, f"status={status}")
        check("预览统计正确（1 有效 / 1 错误 / 模板示例行跳过）",
              data.get("valid_count") == 1 and data.get("error_count") == 1
              and data.get("skipped_examples") == 1,
              f"valid={data.get('valid_count')} error={data.get('error_count')} "
              f"skipped={data.get('skipped_examples')}")
        err_items = [i for i in (data.get("items") or []) if i.get("errors")]
        check("错误行给出原因", bool(err_items) and bool(err_items[0]["errors"]),
              err_items[0]["errors"] if err_items else None)
        check("预览不改库（导入前无该资源）",
              sql(f"select count(*) from emergency_resources where enterprise_id='{ENT}' "
                  f"and name='{NAME_OK}';") == "0")

        ok_item = next(i["data"] for i in (data.get("items") or [])
                       if not i.get("errors") and i["data"]["name"] == NAME_OK)
        check("模板示例行未进入预览", all(
            "【示例】" not in (i["data"].get("name") or "") for i in (data.get("items") or [])))
        status, created = call("POST", f"/enterprises/{ENT}/resources/batch", token,
                               {"items": [ok_item]})
        check("批量落库返回 201", status == 201, f"status={status}")
        check("落库 1 条", len(created.get("data") or []) == 1, created)
        check("资源台账可查到", sql(
            f"select count(*) from emergency_resources where enterprise_id='{ENT}' "
            f"and name='{NAME_OK}';") == "1")
    finally:
        sql(f"delete from emergency_resources where enterprise_id='{ENT}' "
            f"and name in ('{NAME_OK}','{NAME_BAD}');")
        print("cleanup done: 探针资源已删除")

    passed = sum(1 for r in results if r["ok"])
    print(f"合计 {passed}/{len(results)} PASS")
    with open(os.path.join(OUT, "summary-resource-import.json"), "w", encoding="utf-8") as fh:
        json.dump({"passed": passed, "total": len(results), "results": results}, fh,
                  ensure_ascii=False, indent=2)
    return 0 if passed == len(results) else 1


if __name__ == "__main__":
    sys.exit(main())
